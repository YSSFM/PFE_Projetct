import os
import io
import base64
import numpy as np
import cv2
import face_recognition
from datetime import datetime, date, timedelta
from flask import Flask, render_template, request, redirect, url_for, session, send_file
from docx import Document
from scipy.spatial import distance as dist
from database import DatabaseContext
from service.face_service import detect_and_crop_face_cnn, generate_face_embeddings

# Fichier maître regroupant la configuration du serveur Flask et l'aiguillage des requêtes.
app = Flask(__name__)

# Génération d'une clé cryptographique éphémère à chaque initialisation du serveur
app.secret_key = os.urandom(24)

# Variable globale simulant le suivi des tentatives infructueuses de connexion
login_attempts = {}

# FONCTIONS UTILITAIRES POUR L'ANTI-SPOOFING (CLIGNEMENT)
def eye_aspect_ratio(eye):
    # Calcule les distances verticales entre les repères des paupières
    A = dist.euclidean(eye[1], eye[5])
    B = dist.euclidean(eye[2], eye[4])
    # Calcule la distance horizontale entre les coins de l'œil
    C = dist.euclidean(eye[0], eye[3])
    # Calcule le ratio d'aspect de l'œil (EAR)
    ear = (A + B) / (2.0 * C)
    return ear

@app.route('/')
def login_page():
    return render_template('authentification.html')

@app.route('/login', methods=['POST'])
def login():
    username = request.form['username']
    password = request.form['password']

    with DatabaseContext() as (cursor, conn):
        cursor.execute("SELECT * FROM USERS WHERE USERNAME=%s", (username,))
        user_exists = cursor.fetchone()
        cursor.fetchall() 

        cursor.execute("SELECT * FROM USERS WHERE USERNAME=%s AND PASSWORD=%s", (username, password))
        user = cursor.fetchone()

    if user:
        session['username'] = username
        login_attempts[username] = 0
        return redirect(url_for('accueil'))
    else:
        login_attempts[username] = login_attempts.get(username, 0) + 1
        show_reset = False
        if login_attempts[username] >= 3 and user_exists:
            show_reset = True
        return render_template('authentification.html', error=True, show_reset=show_reset)

@app.route('/create_account', methods=['POST'])
def create_account():
    username = request.form['new_username']
    password = request.form['new_password']
    
    with DatabaseContext() as (cursor, conn):
        cursor.execute("INSERT INTO USERS (USERNAME, PASSWORD) VALUES (%s, %s)", (username, password))
        conn.commit()
        
    return redirect(url_for('login_page'))

@app.route('/reset_password', methods=['POST'])
def reset_password():
    username = request.form['reset_username']

    with DatabaseContext() as (cursor, conn):
        cursor.execute("SELECT PASSWORD FROM USERS WHERE USERNAME=%s", (username,))
        result = cursor.fetchone()
        cursor.fetchall() 

    if result:
        recovered_password = result[0]
        return render_template('authentification.html', recovered_user=username, recovered_password=recovered_password)
    else:
        return render_template('authentification.html', show_create=True)

@app.route('/accueil')
def accueil():
    if 'username' not in session:
        return redirect(url_for('login_page'))
    return render_template('index.html', username=session['username'])

# SECTION CRUD ÉTUDIANT
@app.route('/etudiants', methods=['GET'])
def liste_etudiants():
    if 'username' not in session:
        return redirect(url_for('login_page'))
        
    with DatabaseContext() as (cursor, conn):
        cursor.execute("""
            SELECT e.CNE_MASSAR, e.NOM, e.PRENOM, e.SEX, e.TELEPHONE, e.MAIL, e.OPTIONS, e.ANNEE_ACADEMIQUE, e.SEMESTRE 
            FROM etudiant e
        """)
        etudiants = cursor.fetchall()
        
    return render_template('liste_etudiants.html', etudiants=etudiants)

@app.route('/ajouter_etudiant', methods=['GET'])
def detail_etudiant():
    return render_template('detail_etudiant.html')

@app.route('/ajouter', methods=['POST'])
def ajouter():
    cne = request.form['CNEMASSAR']
    nom = request.form['nom']
    prenom = request.form['prenom']
    sex = request.form['sex']
    telephone = request.form['telephone']
    mail = request.form['mail']
    departement = request.form['departement']
    chef_departement = request.form['chef_departement']
    options = request.form['options']
    annee_academique = request.form['annee_academique']
    semestre = request.form['semestre']
    photo_data = request.form.get('photo_data')

    face_detected_successfully = True

    with DatabaseContext() as (cursor, conn):
        # --- SÉCURITÉ : Vérifier si l'étudiant existe déjà ---
        cursor.execute("SELECT CNE_MASSAR FROM etudiant WHERE CNE_MASSAR = %s", (cne,))
        etudiant_existe = cursor.fetchone()
        
        if etudiant_existe:
            # On vide le buffer réseau MySQL par sécurité
            cursor.fetchall()
            # On retourne l'interface avec une variable indiquant que le CNE existe déjà
            return render_template('detail_etudiant.html', success=False, error_cne_double=True, cne_fourni=cne)

        # Si l'étudiant n'existe pas, on continue le traitement normal
        cursor.execute("SELECT ID_DEPARTEMENT FROM departement WHERE NOM_DEPARTEMENT = %s", (departement,))
        dept_res = cursor.fetchone()
        
        if dept_res:
            id_dept = dept_res[0]
        else:
            cursor.execute("INSERT INTO departement (NOM_DEPARTEMENT, CHEF_DEPARTEMENT) VALUES (%s, %s)", (departement, chef_departement))
            id_dept = cursor.lastrowid

        cursor.execute("""
            INSERT INTO etudiant (CNE_MASSAR, ID_DEPARTEMENT, NOM, PRENOM, SEX, TELEPHONE, MAIL, OPTIONS, ANNEE_ACADEMIQUE, SEMESTRE)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (cne, id_dept, nom, prenom, sex, telephone, mail, options, annee_academique, semestre))

        if photo_data and "," in photo_data:
            _, encoded = photo_data.split(",", 1)
            image_bytes = base64.b64decode(encoded)
            
            # Appel sécurisé au service de détection faciale
            face_img, face_bytes = detect_and_crop_face_cnn(image_bytes, use_custom=False)
            
            if face_img is not None:
                os.makedirs("static/photos", exist_ok=True)
                cv2.imwrite(f"static/photos/{cne}.jpg", face_img)
                cursor.execute("INSERT INTO image (CNE_MASSAR, IMAGE) VALUES (%s, %s)", (cne, face_bytes))
            else:
                face_detected_successfully = False

        conn.commit()

    # Si aucun visage n'a pu être extrait de la photo reçue
    if not face_detected_successfully:
        return render_template('detail_etudiant.html', success=False, error_visage=True)

    return render_template('detail_etudiant.html', success=True, CNEMASSAR=cne, nom=nom, prenom=prenom, sex=sex, telephone=telephone, mail=mail, departement=departement, chef_departement=chef_departement, options=options, annee_academique=annee_academique, semestre=semestre)

@app.route('/modifier_etudiant/<cne>', methods=['GET', 'POST'])
def modifier_etudiant(cne):
    if 'username' not in session:
        return redirect(url_for('login_page'))

    with DatabaseContext() as (cursor, conn):
        if request.method == 'POST':
            nom = request.form['nom']
            prenom = request.form['prenom']
            sex = request.form['sex']
            telephone = request.form['telephone']
            mail = request.form['mail']
            options = request.form['options']
            annee = request.form['annee_academique']
            semestre = request.form['semestre']
            photo_data = request.form.get('photo_data')

            cursor.execute("""
                UPDATE etudiant 
                SET NOM=%s, PRENOM=%s, SEX=%s, TELEPHONE=%s, MAIL=%s, OPTIONS=%s, ANNEE_ACADEMIQUE=%s, SEMESTRE=%s 
                WHERE CNE_MASSAR=%s
            """, (nom, prenom, sex, telephone, mail, options, annee, semestre, cne))

            if photo_data and "," in photo_data:
                _, encoded = photo_data.split(",", 1)
                image_bytes = base64.b64decode(encoded)
                face_img, face_bytes = detect_and_crop_face_cnn(image_bytes, use_custom=False)
                
                if face_img is not None:
                    cv2.imwrite(f"static/photos/{cne}.jpg", face_img)
                    cursor.execute("SELECT ID_IMAGE FROM image WHERE CNE_MASSAR=%s", (cne,))
                    exist_img = cursor.fetchone()
                    if exist_img:
                        cursor.execute("UPDATE image SET IMAGE=%s WHERE CNE_MASSAR=%s", (face_bytes, cne))
                    else:
                        cursor.execute("INSERT INTO image (CNE_MASSAR, IMAGE) VALUES (%s, %s)", (cne, face_bytes))

            conn.commit()
            return redirect(url_for('liste_etudiants'))

        cursor.execute("SELECT * FROM etudiant WHERE CNE_MASSAR=%s", (cne,))
        etudiant = cursor.fetchone()
        
    return render_template('modifier_etudiant.html', etudiant=etudiant)

@app.route('/supprimer_etudiant/<cne>', methods=['GET'])
def supprimer_etudiant(cne):
    if 'username' not in session:
        return redirect(url_for('login_page'))

    with DatabaseContext() as (cursor, conn):
        cursor.execute("DELETE FROM presence WHERE CNE_MASSAR=%s", (cne,))
        cursor.execute("DELETE FROM image WHERE CNE_MASSAR=%s", (cne,))
        cursor.execute("DELETE FROM etudiant WHERE CNE_MASSAR=%s", (cne,))
        conn.commit()

    if os.path.exists(f"static/photos/{cne}.jpg"):
        os.remove(f"static/photos/{cne}.jpg")

    return redirect(url_for('liste_etudiants'))

# SECTION PRÉSENCES ET SUIVI DE SÉCURITÉ
@app.route('/presence', methods=['GET', 'POST'])
def presence():
    if 'username' not in session:
        return redirect(url_for('login_page'))

    infos, presences, nb_presence, nb_absence = {}, [], 0, 0
    date_debut, date_fin = '', ''

    with DatabaseContext() as (cursor, conn):
        if request.method == 'GET':
            cursor.execute("""
                SELECT e.CNE_MASSAR, e.NOM, e.PRENOM, d.NOM_DEPARTEMENT, e.MAIL, p.DATE_PRESENCE, p.HEURE, p.STATUT
                FROM etudiant e
                JOIN departement d ON e.ID_DEPARTEMENT = d.ID_DEPARTEMENT
                JOIN presence p ON e.CNE_MASSAR = p.CNE_MASSAR
                ORDER BY p.DATE_PRESENCE DESC
            """)
            presences = cursor.fetchall()

        elif request.method == 'POST':
            cne = request.form.get('cne')
            date_debut = request.form.get('date_debut')
            date_fin = request.form.get('date_fin')
            action = request.form.get('action')

            if cne:
                cursor.execute("SELECT NOM, PRENOM, MAIL FROM etudiant WHERE CNE_MASSAR = %s", (cne,))
                result = cursor.fetchone()
                if result:
                    infos = {'cne': cne, 'nom': result[0], 'prenom': result[1], 'mail': result[2], 'statut': 'Enregistré'}

                if action == 'Calculer' and date_debut and date_fin:
                    cursor.execute("""
                        SELECT e.CNE_MASSAR, e.NOM, e.PRENOM, d.NOM_DEPARTEMENT, e.MAIL, p.DATE_PRESENCE, p.HEURE, p.STATUT
                        FROM etudiant e
                        JOIN departement d ON e.ID_DEPARTEMENT = d.ID_DEPARTEMENT
                        JOIN presence p ON e.CNE_MASSAR = p.CNE_MASSAR
                        WHERE e.CNE_MASSAR = %s AND p.DATE_PRESENCE BETWEEN %s AND %s
                        ORDER BY p.DATE_PRESENCE DESC
                    """, (cne, date_debut, date_fin))
                    presences = cursor.fetchall()

                    start = datetime.strptime(date_debut, "%Y-%m-%d")
                    end = datetime.strptime(date_fin, "%Y-%m-%d")
                    jours_ouvrables = sum(1 for i in range((end - start).days + 1) if (start + timedelta(days=i)).weekday() < 5)

                    nb_presence = sum(1 for p in presences if p[7] == 'Présent')
                    nb_absence = max(0, jours_ouvrables - nb_presence)

                elif action == 'Rapport Étudiant(e)':
                    cursor.execute("""
                        SELECT e.CNE_MASSAR, e.NOM, e.PRENOM, d.NOM_DEPARTEMENT, e.MAIL, p.DATE_PRESENCE, p.HEURE, p.STATUT
                        FROM etudiant e
                        JOIN departement d ON e.ID_DEPARTEMENT = d.ID_DEPARTEMENT
                        JOIN presence p ON e.CNE_MASSAR = p.CNE_MASSAR
                        WHERE e.CNE_MASSAR = %s
                        ORDER BY p.DATE_PRESENCE DESC
                    """, (cne,))
                    presences = cursor.fetchall()

    return render_template('presence.html', infos=infos, presences=presences, nb_presence=nb_presence, nb_absence=nb_absence, date_debut=date_debut, date_fin=date_fin)

@app.route('/rechercher_etudiant', methods=['POST'])
def rechercher_etudiant():
    if 'username' not in session:
        return redirect(url_for('login_page'))

    cne = request.form.get('cne_recherche')
    with DatabaseContext() as (cursor, conn):
        cursor.execute("SELECT NOM, PRENOM, MAIL FROM etudiant WHERE CNE_MASSAR = %s", (cne,))
        result = cursor.fetchone()

    if result:
        infos = {'cne': cne, 'nom': result[0], 'prenom': result[1], 'mail': result[2], 'statut': 'Enregistré'}
        return render_template('presence.html', infos=infos, presences=[], nb_presence=0, nb_absence=0)
    else:
        return render_template('presence.html', popup_cne=cne, infos={}, presences=[], nb_presence=0, nb_absence=0, date_debut='', date_fin='')

@app.route('/entrainement')
def entrainement():
    if 'username' not in session:
        return redirect(url_for('login_page'))
    return render_template('entrainement.html')

@app.route('/lancer_entrainement', methods=['POST'])
def lancer_entrainement():
    try:
        encodings, ids = [], []
        photos_path = "static/photos"
        
        if not os.path.exists(photos_path):
            return render_template('entrainement.html', error=True)

        for filename in os.listdir(photos_path):
            if filename.endswith(".jpg"):
                cne = filename.split(".")[0]
                image_path = os.path.join(photos_path, filename)
                img = cv2.imread(image_path)

                if img is not None:
                    rgb_img = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
                    embedding = generate_face_embeddings(rgb_img, use_custom=False)
                    if embedding is not None:
                        encodings.append(embedding)
                        ids.append(cne)

        np.save("model_encodings.npy", encodings)
        np.save("model_ids.npy", ids)
        return render_template('entrainement.html', success=True)
    except Exception as e:
        print(f"Erreur entrainement : {e}")
        return render_template('entrainement.html', error=True)

@app.route('/detection_faciale')
def detection_faciale():
    if 'username' not in session:
        return redirect(url_for('login_page'))
    return render_template('detection_faciale.html')

# MOTEUR DE RECONNAISSANCE FACIALE + ANTI-SPOOFING (DYNAMIQUE)
@app.route('/scanner', methods=['POST'])
def scanner():
    video = None
    try:
        if not os.path.exists("model_encodings.npy"):
            return render_template('detection_faciale.html', error=True)

        encodings = np.load("model_encodings.npy", allow_pickle=True)
        ids = np.load("model_ids.npy", allow_pickle=True)

        # Initialisation de la caméra avec le backend DirectShow (CAP_DSHOW) recommandé sur Windows
        video = cv2.VideoCapture(0, cv2.CAP_DSHOW)
        
        # Paramètres d'anti-spoofing
        EYE_AR_THRESH = 0.23  
        EYE_AR_CONSEC_FRAMES = 2  
        COUNTER = 0
        TOTAL_BLINKS = 0
        
        cne_detecte = None
        nom_detecte = None
        anti_spoofing_valide = False

        print("Initialisation de la webcam en cours... En attente d'un flux d'images valide.")
        
        # Boucle de pré-chauffage et de stabilisation forcée du capteur
        # On attend d'obtenir au moins une frame qui contient de vrais pixels
        frame_valide_initiale = False
        for _ in range(30):  # Test sur un maximum de 30 tentatives (environ 1 à 2 secondes)
            ret, frame = video.read()
            if ret and frame is not None and frame.size > 0:
                frame_valide_initiale = True
                break
            cv2.waitKey(50) # Petite pause pour laisser le matériel s'activer

        if not frame_valide_initiale:
            print("Erreur critique : Impossible d'obtenir un flux d'images stable de la webcam.")
            if video is not None:
                video.release()
            return render_template('detection_faciale.html', error=True)

        # Boucle principale de traitement vidéo
        while True:
            ret, frame = video.read()
            
            # Si une frame spécifique échoue au milieu du flux, on passe à la suivante sans faire planter l'application
            if not ret or frame is None or frame.size == 0:
                continue

            try:
                # Vérification de la géométrie de la matrice reçue
                if len(frame.shape) < 2:
                    continue
                    
                # Si la webcam capture au format BGRA (4 canaux), éliminer le canal alpha transparent
                if len(frame.shape) == 3 and frame.shape[2] == 4:
                    frame = cv2.cvtColor(frame, cv2.COLOR_BGRA2BGR)
                
                # Conversion explicite vers le format de dlib (RGB au format 8-bit non signé)
                rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                rgb_frame = np.ascontiguousarray(rgb_frame, dtype=np.uint8)
                
            except Exception as e:
                print(f"Échec de conversion de la frame courante : {e}")
                continue

            # Détection des visages sur l'image nettoyée et validée
            face_locations = face_recognition.face_locations(rgb_frame, model="cnn")
            face_encs = face_recognition.face_encodings(rgb_frame, face_locations)
            face_landmarks_list = face_recognition.face_landmarks(rgb_frame, face_locations)

            if not face_locations:
                cv2.putText(frame, "Aucun visage detecte", (20, 40), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 0, 255), 2)

            for (top, right, bottom, left), face_enc, face_landmarks in zip(face_locations, face_encs, face_landmarks_list):
                distances = face_recognition.face_distance(encodings, face_enc)
                
                if len(distances) > 0 and min(distances) < 0.45:
                    index = np.argmin(distances)
                    cne_temp = ids[index]
                    
                    cv2.rectangle(frame, (left, top), (right, bottom), (0, 255, 0), 2)
                    
                    left_eye = face_landmarks['left_eye']
                    right_eye = face_landmarks['right_eye']
                    
                    leftEAR = eye_aspect_ratio(left_eye)
                    rightEAR = eye_aspect_ratio(right_eye)
                    ear = (leftEAR + rightEAR) / 2.0

                    if ear < EYE_AR_THRESH:
                        COUNTER += 1
                    else:
                        if COUNTER >= EYE_AR_CONSEC_FRAMES:
                            TOTAL_BLINKS += 1
                        COUNTER = 0

                    cv2.putText(frame, f"Clignements requis : {TOTAL_BLINKS}/1", (left, top - 35), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 0), 2)
                    cv2.putText(frame, f"EAR: {ear:.2f}", (left, top - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 0), 1)

                    if TOTAL_BLINKS >= 1:
                        cne_detecte = cne_temp
                        anti_spoofing_valide = True
                        break
                else:
                    cv2.rectangle(frame, (left, top), (right, bottom), (0, 0, 255), 2)
                    cv2.putText(frame, "Inconnu - Enregistrer d'abord", (left, top - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 0, 255), 2)

            cv2.imshow("Verification d'identite et anti-spoofing (Echap pour quitter)", frame)
            
            # Sortie de la boucle si l'étudiant cligne des yeux ou si la touche Échap (27) est pressée
            if anti_spoofing_valide or cv2.waitKey(1) == 27:
                break

    except Exception as e:
        print(f"Erreur interne critique au sein du scanner : {e}")
        return render_template('detection_faciale.html', error=True)
        
    finally:
        # Libération systématique de la ressource matérielle pour éviter de bloquer la webcam
        if video is not None:
            video.release()
        cv2.destroyAllWindows()

    # Traitement d'enregistrement de la présence si la détection a réussi
    if anti_spoofing_valide and cne_detecte:
        with DatabaseContext() as (cursor, conn):
            cursor.execute("SELECT NOM FROM etudiant WHERE CNE_MASSAR=%s", (cne_detecte,))
            etudiant = cursor.fetchone()
            
            if etudiant:
                nom_detecte = etudiant[0]
                cursor.execute("SELECT * FROM presence WHERE CNE_MASSAR=%s AND DATE_PRESENCE=%s", (cne_detecte, date.today()))
                deja_present = cursor.fetchone()

                if not deja_present:
                    cursor.execute("INSERT INTO presence (CNE_MASSAR, DATE_PRESENCE, HEURE, STATUT) VALUES (%s, %s, %s, %s)", (cne_detecte, date.today(), datetime.now().time(), "Présent"))
                    conn.commit()
                    return render_template('detection_faciale.html', presence_enregistree=True, nom=nom_detecte, cne=cne_detecte)
                else:
                    return render_template('detection_faciale.html', presence_deja=True, nom=nom_detecte, cne=cne_detecte)

    return render_template('detection_faciale.html', etudiant_inconnu=True)

@app.route('/generer_rapport', methods=['GET', 'POST'])
def generer_rapport():
    with DatabaseContext() as (cursor, conn):
        cursor.execute("""
            SELECT e.CNE_MASSAR, e.NOM, e.PRENOM, e.SEX, e.OPTIONS, d.NOM_DEPARTEMENT, d.CHEF_DEPARTEMENT,
                   e.ANNEE_ACADEMIQUE, e.SEMESTRE, e.MAIL, e.TELEPHONE, p.DATE_PRESENCE, p.HEURE, p.STATUT
            FROM etudiant e
            JOIN departement d ON e.ID_DEPARTEMENT = d.ID_DEPARTEMENT
            JOIN presence p ON e.CNE_MASSAR = p.CNE_MASSAR
            ORDER BY p.DATE_PRESENCE DESC, e.NOM ASC
        """)
        presences = cursor.fetchall()
    return render_template('rapport.html', presences=presences)

@app.route('/imprimer_rapport', methods=['POST'])
def imprimer_rapport():
    with DatabaseContext() as (cursor, conn):
        cursor.execute("""
            SELECT e.CNE_MASSAR, e.NOM, e.PRENOM, e.SEX, e.OPTIONS, d.NOM_DEPARTEMENT, d.CHEF_DEPARTEMENT,
                   e.ANNEE_ACADEMIQUE, e.SEMESTRE, e.MAIL, e.TELEPHONE, p.DATE_PRESENCE, p.HEURE, p.STATUT
            FROM etudiant e
            JOIN departement d ON e.ID_DEPARTEMENT = d.ID_DEPARTEMENT
            JOIN presence p ON e.CNE_MASSAR = p.CNE_MASSAR
            ORDER BY p.DATE_PRESENCE DESC, e.NOM ASC
        """)
        presences = cursor.fetchall()

    doc = Document()
    doc.add_heading('Rapport Général de Présence Étudiante', 0)
    table = doc.add_table(rows=1, cols=7)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'CNE'
    hdr_cells[1].text = 'Nom'
    hdr_cells[2].text = 'Prénom'
    hdr_cells[3].text = 'Filière / Option'
    hdr_cells[4].text = 'Date'
    hdr_cells[5].text = 'Heure'
    hdr_cells[6].text = 'Statut'

    for p in presences:
        row_cells = table.add_row().cells
        row_cells[0].text = str(p[0])
        row_cells[1].text = str(p[1])
        row_cells[2].text = str(p[2])
        row_cells[3].text = str(p[4])
        row_cells[4].text = str(p[11])
        row_cells[5].text = str(p[12])
        row_cells[6].text = str(p[13])

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='rapport_presence.docx')

@app.route('/interface_impression', methods=['POST'])
def interface_impression():
    with DatabaseContext() as (cursor, conn):
        cursor.execute("""
            SELECT e.CNE_MASSAR, e.NOM, e.PRENOM, e.SEX, e.OPTIONS, d.NOM_DEPARTEMENT, d.CHEF_DEPARTEMENT,
                   e.ANNEE_ACADEMIQUE, e.SEMESTRE, e.MAIL, e.TELEPHONE, p.DATE_PRESENCE, p.HEURE, p.STATUT
            FROM etudiant e
            JOIN departement d ON e.ID_DEPARTEMENT = d.ID_DEPARTEMENT
            JOIN presence p ON e.CNE_MASSAR = p.CNE_MASSAR
            ORDER BY p.DATE_PRESENCE DESC, e.NOM ASC
        """)
        presences = cursor.fetchall()
    return render_template('impression.html', presences=presences, now=datetime.now())

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login_page'))

if __name__ == '__main__':
    app.run(debug=True)